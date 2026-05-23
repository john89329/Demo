import os
import zipfile
import tempfile
import shutil
import traceback
from collections import namedtuple

from PyPDF2 import PdfReader
from docx import Document as DocxDocument

ParsedDocument = namedtuple("ParsedDocument", ["file_path", "file_name", "file_type", "content", "pages", "metadata"])


def parse_file(file_path, **kwargs):
    ext = os.path.splitext(file_path)[1].lower()
    parser = PARSER_REGISTRY.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported format: {ext}")
    return parser(file_path, **kwargs)


def parse_docx(file_path, **kwargs):
    doc = DocxDocument(file_path)
    content_parts = []
    pages = []
    current_page = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            current_page.append(text)
        if len("\n".join(current_page)) > 3000:
            pages.append("\n".join(current_page))
            current_page = []
    if current_page:
        pages.append("\n".join(current_page))

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_parts.append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        content_parts.append("\n".join(rows))

    content = "\n".join(content_parts)
    if not pages:
        pages = [content]

    meta = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
    return ParsedDocument(file_path, os.path.basename(file_path), "word", content, pages, meta)


def parse_xlsx(file_path, **kwargs):
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    content_parts = []
    pages = []
    meta = {"sheets": []}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        meta["sheets"].append(sheet_name)
        sheet_content = [f"=== Sheet: {sheet_name} ==="]
        for row in ws.iter_rows(values_only=True):
            row_vals = [str(c) if c is not None else "" for c in row]
            line = "\t".join(row_vals).strip()
            if line:
                sheet_content.append(line)
        sheet_text = "\n".join(sheet_content)
        content_parts.append(sheet_text)
        pages.append(sheet_text)

    wb.close()
    content = "\n\n".join(content_parts)
    return ParsedDocument(file_path, os.path.basename(file_path), "excel", content, pages, meta)


def parse_vsdx(file_path, **kwargs):
    content_parts = []
    pages = []
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            from lxml import etree

            pages_xml_path = None
            for name in z.namelist():
                if name.endswith("pages.xml") and "visio" in name:
                    pages_xml_path = name
                    break

            if pages_xml_path is None:
                for name in z.namelist():
                    if "pages" in name and name.endswith(".xml"):
                        pages_xml_path = name
                        break

            page_ids = []
            if pages_xml_path:
                pages_xml = z.read(pages_xml_path)
                root = etree.fromstring(pages_xml)
                ns = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
                for page_elem in root.iter("{*}Page"):
                    pid = page_elem.get("{*}ID") or page_elem.get("ID")
                    if pid:
                        page_ids.append(pid)

            if not page_ids:
                for name in z.namelist():
                    if "page" in name and name.endswith(".xml") and "pages" not in name:
                        page_ids.append(name)

            for pid in page_ids:
                page_xml_path = None
                for name in z.namelist():
                    if pid in name and name.endswith(".xml") and "pages" not in name:
                        page_xml_path = name
                        break

                if page_xml_path is None:
                    for name in z.namelist():
                        if "page" in name and name.endswith(".xml") and "pages" not in name:
                            page_xml_path = name
                            break

                if page_xml_path and page_xml_path in z.namelist():
                    xml_content = z.read(page_xml_path)
                    root = etree.fromstring(xml_content)
                    texts = []
                    for elem in root.iter():
                        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
                        if tag == "Text" and elem.text and elem.text.strip():
                            texts.append(elem.text.strip())
                    page_text = "\n".join(texts) if texts else ""
                    if page_text:
                        content_parts.append(f"=== Page: {pid} ===\n{page_text}")
                        pages.append(page_text)

    except Exception:
        pass

    content = "\n\n".join(content_parts) if content_parts else f"[Visio diagram: {os.path.basename(file_path)}]"
    return ParsedDocument(file_path, os.path.basename(file_path), "visio", content, pages, {"pages": len(pages)})


def parse_pdf(file_path, skip_ocr=False):
    content_parts = []
    pages = []
    ocr_used = False
    try:
        reader = PdfReader(file_path)
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                pages.append(text)
                content_parts.append(f"--- Page {i + 1} ---\n{text}")
    except Exception:
        pass

    # OCR fallback for scanned PDFs
    if not pages and not skip_ocr:
        try:
            from services.ocr_engine import ocr_pdf
            ocr_text = ocr_pdf(file_path)
            if ocr_text.strip():
                pages = [ocr_text]
                content_parts = [ocr_text]
                ocr_used = True
        except Exception:
            pass

    if not content_parts:
        content_parts.append(f"[PDF could not be parsed: {os.path.basename(file_path)}]")

    content = "\n\n".join(content_parts)
    meta = {"pages": len(pages)}
    if ocr_used:
        meta["ocr_used"] = True
    return ParsedDocument(file_path, os.path.basename(file_path), "pdf", content, pages, meta)


def parse_archive(file_path, **kwargs):
    ext = os.path.splitext(file_path)[1].lower()
    content_parts = []
    all_docs = []
    temp_dir = tempfile.mkdtemp()

    try:
        if ext == ".zip":
            with zipfile.ZipFile(file_path, "r") as z:
                z.extractall(temp_dir)
        elif ext == ".rar":
            try:
                import rarfile
                with rarfile.RarFile(file_path) as rf:
                    rf.extractall(temp_dir)
            except ImportError:
                return ParsedDocument(file_path, os.path.basename(file_path), "archive",
                                      f"[Archive: {os.path.basename(file_path)} (rarfile not available)]",
                                      [], {"error": "rarfile library not available"})

        for root, dirs, files in os.walk(temp_dir):
            for f in files:
                fpath = os.path.join(root, f)
                fext = os.path.splitext(f)[1].lower()
                if fext in PARSER_REGISTRY and fext not in (".zip", ".rar"):
                    try:
                        doc = PARSER_REGISTRY[fext](fpath, **kwargs)
                        all_docs.append(doc)
                        content_parts.append(f"[From archive: {os.path.basename(file_path)}/{f}]\n{doc.content}")
                    except Exception:
                        pass
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    content = "\n\n".join(content_parts) if content_parts else f"[Archive: {os.path.basename(file_path)} (no supported files found)]"
    return ParsedDocument(file_path, os.path.basename(file_path), "archive", content, [content], {"extracted_files": len(all_docs)})


def parse_image(file_path, **kwargs):
    from services.ocr_engine import ocr_image
    text = ocr_image(file_path)
    if not text.strip():
        text = f"[Image: {os.path.basename(file_path)} (no text recognized)]"
    return ParsedDocument(file_path, os.path.basename(file_path), "image", text, [text], {"ocr_used": True})


def parse_xls(file_path, **kwargs):
    import xlrd
    wb = xlrd.open_workbook(file_path)
    content_parts = []
    pages = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        sheet_content = [f"=== Sheet: {sheet_name} ==="]
        for row_idx in range(ws.nrows):
            row_vals = [str(ws.cell_value(row_idx, col_idx)) for col_idx in range(ws.ncols)]
            line = "\t".join(row_vals).strip()
            if line:
                sheet_content.append(line)
        sheet_text = "\n".join(sheet_content)
        content_parts.append(sheet_text)
        pages.append(sheet_text)
    content = "\n\n".join(content_parts) if content_parts else f"[Legacy Excel: {os.path.basename(file_path)}]"
    return ParsedDocument(file_path, os.path.basename(file_path), "excel", content, pages, {"sheets": len(pages)})


def parse_doc(file_path, **kwargs):
    try:
        doc = DocxDocument(file_path)
    except Exception:
        return ParsedDocument(file_path, os.path.basename(file_path), "word",
                              f"[Legacy Word document could not be read: {os.path.basename(file_path)}. Try converting to .docx format.]",
                              [], {"error": "legacy .doc format not supported, convert to .docx"})
    content_parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_parts.append(text)
    content = "\n".join(content_parts) if content_parts else f"[Legacy Word: {os.path.basename(file_path)} (no text extracted)]"
    return ParsedDocument(file_path, os.path.basename(file_path), "word", content, [content], {"legacy": True})


def parse_pptx(file_path, **kwargs):
    from pptx import Presentation
    prs = Presentation(file_path)
    content_parts = []
    pages = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_texts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    slide_texts.append(" | ".join(cells))
        slide_content = "\n".join(slide_texts)
        if slide_content.strip():
            pages.append(slide_content)
            content_parts.append(f"--- Slide {slide_idx + 1} ---\n{slide_content}")
    content = "\n\n".join(content_parts) if content_parts else f"[PowerPoint: {os.path.basename(file_path)} (no text extracted)]"
    return ParsedDocument(file_path, os.path.basename(file_path), "ppt", content, pages, {"slides": len(prs.slides)})


def parse_ppt(file_path, **kwargs):
    try:
        return parse_pptx(file_path)
    except Exception:
        return ParsedDocument(file_path, os.path.basename(file_path), "ppt",
                              f"[Legacy PowerPoint could not be read: {os.path.basename(file_path)}. Try converting to .pptx format.]",
                              [], {"error": "legacy .ppt format not supported, convert to .pptx"})


def parse_dxf(file_path, **kwargs):
    import ezdxf
    try:
        doc = ezdxf.readfile(file_path)
    except Exception:
        return ParsedDocument(file_path, os.path.basename(file_path), "cad",
                              f"[DXF file could not be read: {os.path.basename(file_path)}]",
                              [], {"error": "dxf parse failed"})
    lines = []
    msp = doc.modelspace()
    for entity in msp:
        if entity.dxftype() == "TEXT":
            text = entity.dxf.text.strip()
            if text:
                lines.append(text)
        elif entity.dxftype() == "MTEXT":
            text = entity.text.strip()
            if text:
                lines.append(text)
    content = "\n".join(lines) if lines else f"[DXF CAD file: {os.path.basename(file_path)} (no text entities found)]"
    return ParsedDocument(file_path, os.path.basename(file_path), "cad", content, [content], {"entities": len(lines)})


PARSER_REGISTRY = {
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".vsdx": parse_vsdx,
    ".pdf": parse_pdf,
    ".zip": parse_archive,
    ".rar": parse_archive,
    ".png": parse_image,
    ".jpg": parse_image,
    ".jpeg": parse_image,
    ".bmp": parse_image,
    ".tiff": parse_image,
    ".tif": parse_image,
    ".xls": parse_xls,
    ".doc": parse_doc,
    ".pptx": parse_pptx,
    ".ppt": parse_ppt,
    ".dxf": parse_dxf,
}
